#!/usr/bin/env python3
"""亚铝商城拉取结果（JSON 分片）→ 星选「供应商产品信息收集表」格式 xlsx。

用法：python3 scripts/yalv_to_xlsx.py <含 chunk_*.json 的目录> <out.xlsx>
（浏览器 hook 下载的 yalv_products.json 直接改名成 chunk_000.json 放进目录即可；JSON 含供货价，不进仓库）
chunks_dir 里是 chunk_*.json，每个是数组，元素字段见 pick()（浏览器里精简过的对象）。
"""
import glob
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 亚铝一级分类 → 星选 13 大类（品类列可在商品库审核时再改）
CAT_MAP = {
    'DELIXI德力西': '电工电气',
    'BOSI波斯工具': '工具耗材',
    'Bridge大桥焊材': '工具耗材',
    'Welding焊机及配件': '机械设备',
    'Lifting吊具': '五金紧固',
    'Sabuk&Jaring安全带/网': '安全防护',
    'Webbing Sling吊带': '五金紧固',
    'Safety劳保产品': '安全防护',
    'Mesin施工机械': '机械设备',
    '货架': '日杂用品',
    '国内及本地采购': '日杂用品',
    'MCB人民电器开关': '电工电气',
    'Tangga铝合金爬梯': '工具耗材',
    '脚手架': '建筑钢材',
    'Alat survei测绘工具': '工具耗材',
    'Kawat Besi铁丝': '建筑钢材',
    'Pompa奥利水泵': '给水排水',
    'Lampu照明灯具': '卫浴照明',
    'Sanitary箭牌家居': '卫浴照明',
    'Scaffolding扣件': '建筑钢材',
    'Baut紧固件10月': '五金紧固',
    'Baut紧固件': '五金紧固',
    'Pengencang通丝': '五金紧固',
    'Plywwod建筑模板': '装饰材料',
    '德高树脂瓦': '防水保温',
    '油漆化工类': '装饰材料',
    '东成工具': '工具耗材',
    'Bahan Abrasif磨具': '工具耗材',
    '钢丝软管': '给水排水',
    '托盘物流箱周转筐': '日杂用品',
    '金桥焊材': '工具耗材',
    'SOMY水泵': '给水排水',
}

UNIT_MAP = {
    'PCS': '个', 'PC': '个', 'KG': '公斤', 'M': '米', 'SET': '套', 'BOX': '箱', 'ROLL': '卷',
    'BAG': '袋', 'PACK': '包', 'TON': '吨', 'L': '升', 'SHEET': '张', 'M2': '平方米',
}
UNITS = ['件', '个', '套', '箱', '袋', '包', '桶', '卷', '根', '支', '张', '块', '片', '米', '平方米', '立方米', '吨', '公斤', '升']

# 与 scripts/gen_product_template.py 的 25 列一一对应（表头文本必须含 vendor.php aliasMap 的别名）
HEADERS = [
    '产品名称\nNama Produk *', '品类\nKategori *', '子类 / 用途\nSub-kategori', '品牌\nMerek', '型号\nModel / Tipe',
    '规格参数\nSpesifikasi', '材质\nMaterial', '颜色\nWarna', '计量单位\nSatuan *',
    '供货价\nHarga (IDR) *', '市场参考价\nHarga Pasar', '起订量\nMOQ', '现货状态\nStok', '可供数量\nJumlah Tersedia',
    '交货周期\nWaktu Pengiriman', '运费说明\nOngkir',
    '包装规格\nKemasan', '单件重量\nBerat', '单件尺寸\nDimensi', '产地\nAsal Produk',
    '执行标准 / 认证\nSertifikasi', '质保期\nGaransi', '售后服务\nLayanan Purna Jual',
    '图片链接\nLink Gambar', '产品描述\nDeskripsi',
]


def unit_of(raw: str) -> str:
    raw = (raw or '').strip()
    if not raw:
        return '件'
    cn = raw.split('/')[0].strip()
    if cn in UNITS:
        return cn
    up = cn.upper()
    if up in UNIT_MAP:
        return UNIT_MAP[up]
    for part in raw.split('/'):
        p = part.strip()
        if p in UNITS:
            return p
        if p.upper() in UNIT_MAP:
            return UNIT_MAP[p.upper()]
    return cn or '件'


def split_name(name: str, spec: str):
    """「漏电断路器/CDM3LS-125C/4300B 125A」→ 名称 + 规格。名称取第一个「/」前；规格优先用 specNo。"""
    name = (name or '').strip()
    spec = (spec or '').strip()
    if '/' in name:
        head, tail = name.split('/', 1)
        head, tail = head.strip(), tail.strip()
        if head:
            return head, (spec or tail)
    return name, spec


def num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f


# 页面里 window.__pull.compact() 的定位数组列序（改了 compact 必须同步改这里）
COMPACT_KEYS = ['id', 'code', 'name', 'spec', 'brand', 'img', 'uom', 'moq', 'cat', 'price', 'retail', 'avail', 'soldOut', 'desc']


def norm(r):
    """兼容两种输入：定位数组（页面 compact()）或字典。"""
    if isinstance(r, list):
        return dict(zip(COMPACT_KEYS, r))
    return r


def main():
    src, out = sys.argv[1], sys.argv[2]
    rows = []
    for f in sorted(glob.glob(str(Path(src) / 'chunk_*.json'))):
        rows.extend(json.load(open(f, encoding='utf-8')))
    seen = set()
    items = []
    for r in rows:
        r = norm(r)
        if r['id'] in seen:
            continue
        seen.add(r['id'])
        items.append(r)

    wb = Workbook()
    ws = wb.active
    ws.title = '产品清单 Produk'
    ws_np = wb.create_sheet('无供货价_不导入')
    for sheet in (ws, ws_np):
        sheet.append(['亚铝商城拉取 · 自动生成'] + [''] * (len(HEADERS) - 1))
        sheet.append(HEADERS)
        for i, h in enumerate(HEADERS, start=1):
            c = sheet.cell(row=2, column=i)
            c.font = Font(bold=True, size=10)
            c.fill = PatternFill('solid', fgColor='F2F4F7')
            c.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.column_dimensions[get_column_letter(i)].width = 16
        sheet.row_dimensions[2].height = 34
        sheet.freeze_panes = 'B3'

    stats = Counter()
    cat_unknown = Counter()
    for it in items:
        cat_raw = (it.get('cat') or '').rstrip('^')
        cat = CAT_MAP.get(cat_raw)
        if not cat:
            cat_unknown[cat_raw] += 1
            cat = '日杂用品'
        name, spec = split_name(it.get('name'), it.get('spec'))
        price = num(it.get('price'))
        retail = num(it.get('retail'))
        avail = num(it.get('avail'))
        moq = num(it.get('moq'))
        desc_parts = []
        if it.get('desc'):
            desc_parts.append(str(it['desc']).strip())
        desc_parts.append(f'亚铝商城编码 {it.get("code", "")}，一级分类「{cat_raw}」')
        row = [
            name, cat, cat_raw, it.get('brand') or '', it.get('code') or '',
            spec, '', '', unit_of(it.get('uom')),
            price if price and price > 0 else '', retail if retail and retail > 0 else '',
            int(moq) if moq and moq == int(moq) else (moq or ''),
            '需订货' if it.get('soldOut') else '有现货',
            int(avail) if avail is not None and avail >= 0 and avail == int(avail) else (avail if avail is not None and avail >= 0 else ''),
            '', '', '', '', '', '', '', '', '',
            it.get('img') or '', '；'.join(desc_parts),
        ]
        if price and price > 0:
            ws.append(row)
            stats['ok'] += 1
        else:
            ws_np.append(row)
            stats['no_price'] += 1
        stats[cat] += 1

    wb.save(out)
    print(f'商品 {len(items)} 条：可导入 {stats["ok"]}，无供货价 {stats["no_price"]} → {out}')
    if cat_unknown:
        print('未映射分类：', dict(cat_unknown))
    print('品类分布：', {k: v for k, v in stats.items() if k not in ('ok', 'no_price')})


if __name__ == '__main__':
    main()
